import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank_layout)

txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(11.333), Inches(2))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = "Test PPTX Widescreen"
p.font.size = Pt(32)
p.font.bold = True
p.font.color.rgb = RGBColor(30, 58, 138)

prs.save("scratch/test_out.pptx")
print("Saved test pptx successfully!")
